from pathlib import Path

from langchain_core.documents import Document
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}


def _base_metadata(path: Path) -> dict[str, str]:
    return {
        "source": str(path),
        "filename": path.name,
    }


def _load_pdf(path: Path) -> list[Document]:
    reader = PdfReader(str(path))
    pages: list[Document] = []
    for index, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if not text.strip():
            continue
        pages.append(
            Document(
                page_content=text,
                metadata={**_base_metadata(path), "page": str(index + 1)},
            )
        )
    if not pages:
        raise ValueError(f"No extractable text found in PDF: {path.name}")
    return pages


def _load_docx(path: Path) -> list[Document]:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    paragraphs = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
    text = "\n\n".join(paragraphs)
    if not text.strip():
        raise ValueError(f"No extractable text found in DOCX: {path.name}")
    return [Document(page_content=text, metadata=_base_metadata(path))]


def _load_text(path: Path) -> list[Document]:
    text = path.read_text(encoding="utf-8")
    if not text.strip():
        raise ValueError(f"File is empty: {path.name}")
    return [Document(page_content=text, metadata=_base_metadata(path))]


def load_document(path: str | Path) -> list[Document]:
    """Load a document from disk and return LangChain Document objects."""
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported file type '{suffix}'. Supported: {supported}")

    if suffix == ".pdf":
        return _load_pdf(path)
    if suffix == ".docx":
        return _load_docx(path)
    return _load_text(path)
